"""Render spec/kommentarliste_pflichtenheft_v1_0.yaml as the customer's
comment list: docs/Kommentarliste_Pflichtenheft_V1.0.xlsx.

One visible sheet "Kommentare": a title line, then the table Nr. | ID |
Kommentar | Bewertung | Antwort Lieferant, the Bewertung column a drop-down
over the values of a hidden sheet "Listen".

The output is reproducible byte for byte: the workbook's timestamps (the
document properties and every ZIP entry) are fixed, and the entries are
stored uncompressed, so the bytes depend neither on the time of the build
nor on the zlib build of the Python running it:

    pip install -r requirements-dev.txt
    python tools/build_comment_list_xlsx.py
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / "spec" / "kommentarliste_pflichtenheft_v1_0.yaml"
OUT = ROOT / "docs" / "Kommentarliste_Pflichtenheft_V1.0.xlsx"


def _reproducible(data: bytes, stamp: datetime) -> bytes:
    """The workbook with fixed timestamps, its entries stored in order."""
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as target,
    ):
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)",
                    rb"\g<1>" + iso + rb"\g<2>",
                    content,
                )
            entry = zipfile.ZipInfo(info.filename, date_time=stamp.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_STORED
            entry.external_attr = 0o600 << 16
            entry.create_system = 0
            target.writestr(entry, content)
    return out.getvalue()


def build() -> bytes:
    spec = yaml.safe_load(SPEC.read_text(encoding="utf-8"))
    stamp = spec["timestamp"]
    if isinstance(stamp, str):
        stamp = datetime.fromisoformat(stamp.replace("Z", "+00:00"))
    stamp = stamp.replace(tzinfo=None)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = spec["sheet"]
    sheet.append([spec["title"]])
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([spec["subtitle"]])
    sheet.append([])
    sheet.append(spec["header"])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for row in spec["rows"]:
        sheet.append([value if value != "" else None for value in row])
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=header_row, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = wrap
    for index, width in enumerate(spec["widths"]):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"

    choices = book.create_sheet("Listen")
    choices.append(["Bewertung"])
    for value in spec["choices"]:
        choices.append([value])
    choices.sheet_state = "hidden"
    column = chr(ord("A") + spec["header"].index("Bewertung"))
    validation = DataValidation(
        type="list",
        formula1=f"=Listen!$A$2:$A${len(spec['choices']) + 1}",
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{column}{header_row + 1}:{column}{sheet.max_row}")

    book.properties.creator = spec["author"]
    book.properties.lastModifiedBy = spec["author"]
    book.properties.title = spec["title"]
    book.properties.created = stamp
    book.properties.modified = stamp
    buffer = io.BytesIO()
    book.save(buffer)
    return _reproducible(buffer.getvalue(), stamp)


if __name__ == "__main__":
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_bytes(build())
    print(f"wrote {OUT.relative_to(ROOT)}")
